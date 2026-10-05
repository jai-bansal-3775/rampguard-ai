import docx
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls


def create_brd_document():
  doc = docx.Document()

  # Helper functions
  def add_header(text, level):
    return doc.add_heading(text, level=level)

  def add_quote(text):
    p = doc.add_paragraph()
    p.style = 'Quote'
    run = p.add_run(text)
    run.italic = True
    return p

  def add_table(headers, data):
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = 'Table Grid'
    hdr_cells = table.rows[0].cells
    for i, header_text in enumerate(headers):
      hdr_cells[i].text = header_text
      # Bold header text
      for paragraph in hdr_cells[i].paragraphs:
        for run in paragraph.runs:
          run.bold = True
      # Set header background color
      shading = parse_xml(r'<w:shd {} w:fill="ECECEC"/>'.format(nsdecls('w')))
      hdr_cells[i]._tc.get_or_add_tcPr().append(shading)

    for row_data in data:
      row_cells = table.add_row().cells
      for i, cell_value in enumerate(row_data):
        row_cells[i].text = cell_value
    doc.add_paragraph()  # Spacer

  # Document Title
  title = add_header('RAMPGUARD AI', level=0)
  p_sub = doc.add_paragraph('Ground Safety and Handling Monitor\n')
  p_sub.runs[0].bold = True

  add_header('BUSINESS REQUIREMENTS DOCUMENT', level=1)

  add_quote(
      'Vision: An AI-assisted safety supervisor that observes aircraft'
      ' turnaround activity, identifies hazardous conditions, explains the'
      ' operational risk, and enables authorized ground teams to intervene'
      ' before an incident occurs.'
  )

  add_table(
      ['Document field', 'Value'],
      [
          ['Document version', '1.0'],
          ['Prepared for', 'Google Cloud AI Builder Cup 2026'],
          [
              'Recommended theme',
              'Manufacturing: Intelligent Operations & Industrial Efficiency',
          ],
          [
              'Prototype scope',
              'One gate, selected safety scenarios, near-real-time monitoring',
          ],
          ['Document date', '3 October 2026'],
      ],
  )

  doc.add_paragraph(
      'Prepared as a competition-aligned working BRD. Operational procedures,'
      ' thresholds, and alert ownership require validation by qualified airport'
      ' and airline safety stakeholders before real-world use.'
  )

  add_header('Document Control', level=2)
  add_table(
      ['Field', 'Value'],
      [
          ['Document owner', 'RampGuard AI project team'],
          ['Business owner', 'To be assigned'],
          ['Product owner', 'To be assigned'],
          ['Status', 'Draft for prototype planning'],
          [
              'Approval authority',
              'Prototype sponsor and qualified ground-safety representative',
          ],
          ['Review cadence', 'At each prototype milestone'],
      ],
  )

  add_header('Revision History', level=2)
  add_table(
      ['Version', 'Date', 'Change', 'Author'],
      [[
          '1.0',
          '3 Oct 2026',
          'Initial BRD aligned to published AI Builder Cup 2026 requirements',
          'RampGuard AI team',
      ]],
  )

  add_header('1. Executive Summary', level=2)
  doc.add_paragraph(
      'A focused, explainable, human-in-the-loop safety prototype for aircraft'
      ' ground operations. RampGuard AI is a proposed ground safety and'
      ' handling monitor for airport ramp and gate environments. The prototype'
      ' will analyze authorized camera or recorded video feeds, detect'
      ' selected safety-relevant conditions, combine visual observations with'
      ' turnaround context, assign a transparent risk level, and present'
      ' actionable alerts to an authorized supervisor.'
  )
  doc.add_paragraph(
      'The business need is to improve the speed and consistency with which'
      ' ground teams become aware of hazardous conditions during aircraft'
      ' turnaround. RampGuard AI is not intended to replace trained'
      ' personnel, certify compliance, command equipment, or autonomously'
      ' stop an operation. It is a decision-support layer that preserves human'
      ' authority and creates an auditable evidence trail.'
  )

  add_quote(
      'Prototype value proposition: Observe -> understand context -> assess'
      ' risk -> alert a human -> capture disposition. The differentiator is'
      ' contextual multimodal reasoning, not camera detection alone.'
  )

  add_header('Recommended competition positioning', level=3)
  doc.add_paragraph(
      '• Primary challenge theme: Manufacturing: Intelligent Operations &'
      ' Industrial Efficiency, positioned around safety, reliability,'
      ' operational efficiency, and turning operational data into'
      ' recommendations or actions.'
  )
  doc.add_paragraph(
      '• Secondary narrative: Future of Work & Enterprise Productivity,'
      ' because the solution supports frontline decision-making and workflow'
      ' execution.'
  )
  doc.add_paragraph(
      '• Prototype requirement: demonstrate an end-to-end functional workflow,'
      ' integrate Google AI, and deploy on Google Cloud using an allowed'
      ' deployment path.'
  )

  add_header('2. Competition Alignment', level=2)
  add_table(
      ['Published criterion', 'Weight', 'RampGuard response'],
      [
          [
              'Technical Merit & Gen AI Implementation',
              '40%',
              (
                  'Gemini-based multimodal interpretation and contextual risk'
                  ' reasoning; working end-to-end prototype; Cloud Run or'
                  ' Firebase deployment; modular event architecture.'
              ),
          ],
          [
              'Problem Alignment & Impact',
              '25%',
              (
                  'Addresses meaningful safety and efficiency challenges in'
                  ' aviation ground operations with clear users, workflow'
                  ' integration, and measurable outcomes.'
              ),
          ],
          [
              'Innovation & Creativity',
              '25%',
              (
                  'Combines visual events, turnaround phase, rules, evidence,'
                  ' and natural-language explanations instead of providing a'
                  ' basic chatbot or isolated object detector.'
              ),
          ],
          [
              'User Experience & Solution Design',
              '10%',
              (
                  'Supervisor-first console with prioritized alerts, evidence,'
                  ' recommended response, acknowledgment, accessibility, and'
                  ' minimal interaction burden.'
              ),
          ],
      ],
  )

  add_header('Mandatory submission alignment', level=3)
  add_table(
      ['Requirement', 'BRD response / deliverable'],
      [
          [
              'Clear proposal in deck/PDF',
              (
                  'Problem, impact, feasibility, scalability, and competition'
                  ' alignment are defined here and can be distilled into the'
                  ' submission deck.'
              ),
          ],
          [
              'Functional prototype',
              (
                  'MVP requirements define a working scenario from video'
                  ' input through detection, reasoning, alerting,'
                  ' acknowledgment, and incident log.'
              ),
          ],
          [
              'Google AI and Google Cloud',
              (
                  'Solution requirements specify Gemini or another permitted'
                  ' Google AI model and deployment on Cloud Run or Firebase/GCP.'
              ),
          ],
          [
              'Three-minute public demo link',
              (
                  'Section 12 provides a timed demonstration storyline for'
                  ' an English-language video.'
              ),
          ],
          [
              'English submission materials',
              'This BRD and proposed submission artifacts are in English.',
          ],
          [
              'User-friendly, accessible, scalable',
              (
                  'Non-functional and UX requirements establish accessibility,'
                  ' modularity, performance targets, and expansion patterns.'
              ),
          ],
      ],
  )

  add_quote(
      'Scope discipline: For the hackathon, depth is more valuable than'
      ' breadth. The MVP should prove two or three safety scenarios reliably in'
      ' one gate environment rather than claim universal ramp coverage.'
  )

  add_header('3. Business Context and Problem', level=2)
  add_header('3.1 Business context', level=3)
  doc.add_paragraph(
      'Aircraft turnaround is a tightly coordinated sequence involving'
      ' people, vehicles, ground support equipment, aircraft zones, and'
      ' time-sensitive activities. Safety observations can be fragmented'
      ' across camera views, local procedures, radio communication, and human'
      ' attention. A prototype that creates a common, explainable risk picture'
      ' can improve situational awareness without taking authority away from'
      ' trained operators.'
  )

  add_header('3.2 Problem statement', level=3)
  add_quote(
      'Problem statement: Ground supervisors may not receive timely,'
      ' prioritized, and context-rich visibility of emerging hazards around an'
      ' aircraft. Conventional monitoring can identify an object or motion but'
      ' may not explain why it matters at the current turnaround phase, what'
      ' evidence supports the alert, or what authorized human action should'
      ' follow.'
  )

  add_header('3.3 Target business outcomes', level=3)
  doc.add_paragraph(
      '• Shorten the time between a visible hazardous condition and supervisor'
      ' awareness.'
  )
  doc.add_paragraph(
      '• Improve consistency of risk classification and alert information.'
  )
  doc.add_paragraph(
      '• Reduce alert fatigue by correlating visual events with operational'
      ' context.'
  )
  doc.add_paragraph(
      '• Create traceable records for review, learning, and process'
      ' improvement.'
  )
  doc.add_paragraph(
      '• Enable a scalable architecture that can add cameras, gates, event'
      ' types, and integrations after validation.'
  )

  add_header('3.4 Design principles', level=3)
  add_table(
      ['Principle', 'Application'],
      [
          [
              'Safety first',
              (
                  'When uncertain, the system communicates uncertainty and'
                  ' defers to authorized personnel.'
              ),
          ],
          [
              'Human authority',
              (
                  'No autonomous operational command, equipment control, or'
                  ' compliance certification.'
              ),
          ],
          [
              'Evidence before assertion',
              (
                  'Every alert includes source evidence, event time, confidence,'
                  ' and reasoning summary.'
              ),
          ],
          [
              'Privacy by design',
              (
                  'Process only authorized feeds; minimize identity data; apply'
                  ' retention and access controls.'
              ),
          ],
          [
              'Practical prototype',
              (
                  'Use a small, demonstrable scenario set and synthetic or'
                  ' approved data.'
              ),
          ],
          [
              'Explainable Gen AI',
              (
                  'Gemini adds context and explanations; deterministic controls'
                  ' guard critical alert logic.'
              ),
          ],
      ],
  )

  add_header('4. Vision, Objectives and Scope', level=2)
  add_header('4.1 Product vision', level=3)
  doc.add_paragraph(
      'For ground-operation supervisors who need rapid awareness of ramp'
      ' hazards, RampGuard AI is a multimodal safety monitoring and'
      ' decision-support solution that turns authorized video and turnaround'
      ' context into prioritized, explainable alerts. Unlike an isolated'
      ' computer-vision detector, it reasons over the operational situation,'
      ' presents supporting evidence, and requires human disposition.'
  )

  add_header('4.2 Objectives', level=3)
  add_table(
      ['ID', 'Objective', 'Prototype evidence'],
      [
          [
              'OBJ-01',
              (
                  'Detect selected observable safety conditions in an'
                  ' approved video stream or clip.'
              ),
              'Annotated event with timestamp and confidence.',
          ],
          [
              'OBJ-02',
              (
                  'Correlate the observation with aircraft zone and turnaround'
                  ' phase.'
              ),
              'Context panel and reasoning trace.',
          ],
          [
              'OBJ-03',
              'Prioritize alerts by severity and urgency.',
              'Ordered alert queue with risk rationale.',
          ],
          [
              'OBJ-04',
              (
                  'Provide a clear, approved-response recommendation to a'
                  ' supervisor.'
              ),
              (
                  'Action card with acknowledge/escalate/false-positive'
                  ' options.'
              ),
          ],
          [
              'OBJ-05',
              'Maintain an auditable event and disposition record.',
              'Incident timeline and exportable record.',
          ],
          [
              'OBJ-06',
              'Demonstrate a scalable Google Cloud architecture.',
              'Deployed prototype and architecture view.',
          ],
      ],
  )

  add_header('4.3 MVP in scope', level=3)
  doc.add_paragraph(
      '• One gate or stand represented by recorded or simulated video.'
  )
  doc.add_paragraph(
      '• Two to three selected scenarios, recommended: personnel/equipment in'
      ' a configured safety zone, PPE non-compliance where visually'
      ' assessable, and FOD-like object detection.'
  )
  doc.add_paragraph('• Configurable zones and turnaround phase.')
  doc.add_paragraph(
      '• Multimodal event interpretation and natural-language explanation using'
      ' a permitted Google AI model.'
  )
  doc.add_paragraph(
      '• Risk classification, supervisor alert, acknowledgment, escalation,'
      ' false-positive feedback, and event log.'
  )
  doc.add_paragraph(
      '• Web console deployed on Google Cloud through Cloud Run or Firebase.'
  )
  doc.add_paragraph('• Demo dataset, evaluation set, and basic metrics dashboard.')

  add_header('4.4 Out of scope for prototype', level=3)
  doc.add_paragraph(
      '• Live control of aircraft, tugs, jet bridges, vehicles, alarms, or'
      ' airport systems.'
  )
  doc.add_paragraph(
      '• Automated stop/go clearance or regulatory compliance determination.'
  )
  doc.add_paragraph(
      '• Biometric identification, facial recognition, or workforce'
      ' performance scoring.'
  )
  doc.add_paragraph(
      '• Full airport deployment, multi-camera tracking across the entire'
      ' apron, or production SLAs.'
  )
  doc.add_paragraph(
      '• Prediction of maintenance failures or passenger-facing functions.'
  )
  doc.add_paragraph(
      '• Use of unapproved operational, employee, customer, or'
      ' security-sensitive data.'
  )

  add_header('5. Stakeholders and Users', level=2)
  add_table(
      ['Stakeholder / user', 'Need', 'RampGuard interaction'],
      [
          [
              'Ground operations supervisor',
              'A consolidated, prioritized view of safety-relevant events.',
              'Reviews alert, evidence, recommendation, and disposition.',
          ],
          [
              'Ramp safety officer',
              'Consistent records for trend review and procedure improvement.',
              'Reviews incident history and alert quality.',
          ],
          [
              'Ground handling team lead',
              'Timely, unambiguous information for local action.',
              'Receives supervisor-authorized escalation.',
          ],
          [
              'Airport / airline operations control',
              'Awareness of material operational risk and disruption.',
              'Consumes approved escalations and summary status.',
          ],
          [
              'System administrator',
              (
                  'Controlled configuration, users, zones, integrations, and'
                  ' retention.'
              ),
              'Administers policies and monitors health.',
          ],
          [
              'Data / AI steward',
              'Model quality, drift, data lineage, and responsible-AI controls.',
              'Reviews evaluations, versioning, and feedback.',
          ],
          [
              'Hackathon evaluator',
              (
                  'Clear proof of technical merit, impact, innovation, and'
                  ' usability.'
              ),
              'Experiences the three-minute end-to-end demo.',
          ],
      ],
  )

  add_header('User persona: shift supervisor', level=3)
  add_quote(
      'Primary persona: Needs a low-noise view that answers four questions'
      ' immediately: What happened? Where and when? Why is it risky now? What'
      ' approved response should I consider?'
  )

  add_header('6. Future-State Business Process', level=2)
  doc.add_paragraph(
      '1. Authorized video or a demo clip enters the monitoring pipeline.'
  )
  doc.add_paragraph(
      '2. Vision processing detects relevant entities, zones, motion, and'
      ' candidate events.'
  )
  doc.add_paragraph(
      '3. Turnaround context service provides gate, phase, and configured'
      ' safety rules.'
  )
  doc.add_paragraph(
      '4. Gemini interprets the candidate event and produces a structured'
      ' explanation with uncertainty explicitly represented.'
  )
  doc.add_paragraph(
      '5. Risk orchestration applies deterministic guardrails and produces the'
      ' alert priority.'
  )
  doc.add_paragraph(
      '6. The supervisor console displays evidence, context, recommendation,'
      ' and confidence.'
  )
  doc.add_paragraph(
      '7. The supervisor acknowledges, escalates, dismisses, or marks a false'
      ' positive.'
  )
  doc.add_paragraph(
      '8. The event, decision, model/configuration version, and feedback are'
      ' logged for audit and improvement.'
  )

  add_header('Alert lifecycle', level=3)
  add_table(
      ['State', 'Meaning', 'Allowed next states'],
      [
          [
              'Detected',
              'Candidate event created by sensor/vision layer.',
              'Validated, Suppressed',
          ],
          [
              'Validated',
              'Context and guardrails confirm alert eligibility.',
              'Presented, Suppressed',
          ],
          [
              'Presented',
              'Visible to authorized user.',
              'Acknowledged, Escalated, Dismissed',
          ],
          ['Acknowledged', 'User has reviewed the alert.', 'Resolved, Escalated'],
          [
              'Escalated',
              'Forwarded through an approved human workflow.',
              'Resolved',
          ],
          [
              'Dismissed',
              'Determined non-actionable or false positive.',
              'Closed',
          ],
          ['Resolved', 'Hazard or concern addressed.', 'Closed'],
          [
              'Closed',
              'Record complete and immutable except approved annotations.',
              'None',
          ],
      ],
  )

  add_header('7. Business and Functional Requirements', level=2)
  add_header('7.1 Business requirements', level=3)
  add_table(
      ['ID', 'Requirement', 'Priority'],
      [
          [
              'BR-01',
              (
                  'The solution shall improve situational awareness for'
                  ' selected ground safety hazards without replacing the'
                  ' authority of trained personnel.'
              ),
              'Must',
          ],
          [
              'BR-02',
              (
                  'The prototype shall demonstrate an end-to-end flow using a'
                  ' working deployment on an allowed Google Cloud path.'
              ),
              'Must',
          ],
          [
              'BR-03',
              (
                  'Every actionable alert shall include evidence, context, risk'
                  ' rationale, and a human disposition path.'
              ),
              'Must',
          ],
          [
              'BR-04',
              (
                  'The solution shall support configuration for gate zones,'
                  ' severity rules, and enabled scenarios.'
              ),
              'Must',
          ],
          [
              'BR-05',
              (
                  'The prototype shall use only synthetic, public, or'
                  ' explicitly approved data.'
              ),
              'Must',
          ],
          [
              'BR-06',
              (
                  'The architecture shall support future expansion to'
                  ' additional event types, feeds, and integrations.'
              ),
              'Should',
          ],
          [
              'BR-07',
              (
                  'The project shall provide measurable evaluation results'
                  ' and disclose prototype limitations.'
              ),
              'Must',
          ],
      ],
  )

  add_header('7.2 Functional requirements', level=3)
  add_table(
      ['ID', 'Functional requirement', 'Priority'],
      [
          [
              'FR-01',
              (
                  'Ingest recorded video, image sequences, or an authorized'
                  ' live feed through a documented interface.'
              ),
              'Must',
          ],
          [
              'FR-02',
              (
                  'Allow an administrator to configure one or more polygonal'
                  ' or logical safety zones.'
              ),
              'Must',
          ],
          [
              'FR-03',
              (
                  'Detect configured classes relevant to the selected scenario,'
                  ' such as person, vehicle, ground equipment, or foreign'
                  ' object candidate.'
              ),
              'Must',
          ],
          [
              'FR-04',
              (
                  'Create a candidate event when configured spatial, temporal,'
                  ' or PPE conditions are met.'
              ),
              'Must',
          ],
          [
              'FR-05',
              (
                  'Associate an event with gate/stand, camera, timestamp,'
                  ' turnaround phase, and scenario type.'
              ),
              'Must',
          ],
          [
              'FR-06',
              (
                  'Use Gemini or another permitted Google AI model to produce'
                  ' a structured event interpretation and concise explanation.'
              ),
              'Must',
          ],
          [
              'FR-07',
              (
                  'Apply deterministic rules to constrain severity and prevent'
                  ' a generative model from independently issuing operational'
                  ' commands.'
              ),
              'Must',
          ],
          [
              'FR-08',
              (
                  'Display alert severity, evidence frame or clip, confidence,'
                  ' reasoning, location, time, and recommended human response.'
              ),
              'Must',
          ],
          [
              'FR-09',
              (
                  'Allow authorized users to acknowledge, escalate, dismiss,'
                  ' resolve, and add a note.'
              ),
              'Must',
          ],
          [
              'FR-10',
              (
                  'Record user disposition, timestamps, model version, rule'
                  ' version, and evidence reference.'
              ),
              'Must',
          ],
          [
              'FR-11',
              (
                  'Suppress duplicate alerts within a configurable period and'
                  ' display suppression status.'
              ),
              'Should',
          ],
          [
              'FR-12',
              (
                  'Provide search and filtering by time, gate, scenario,'
                  ' priority, and status.'
              ),
              'Should',
          ],
          [
              'FR-13',
              (
                  'Provide a basic dashboard for alert volume, disposition,'
                  ' false-positive feedback, and response time.'
              ),
              'Should',
          ],
          [
              'FR-14',
              (
                  'Provide an administrator view for scenario enablement,'
                  ' thresholds, roles, and retention configuration.'
              ),
              'Could',
          ],
          [
              'FR-15',
              (
                  'Provide a demo mode that replays a curated incident sequence'
                  ' predictably.'
              ),
              'Must',
          ],
      ],
  )

  add_header('7.3 Core use cases', level=3)
  add_table(
      ['Use case', 'Trigger', 'System behavior', 'Human outcome'],
      [
          [
              'UC-01 Zone intrusion',
              (
                  'Person or equipment enters a configured zone during an'
                  ' incompatible phase.'
              ),
              (
                  'Correlates location and phase, presents evidence and risk'
                  ' rationale.'
              ),
              'Supervisor verifies and initiates approved local response.',
          ],
          [
              'UC-02 PPE observation',
              (
                  'Required visible PPE is not detected for a person in a'
                  ' configured area.'
              ),
              'Shows confidence and avoids identity inference.',
              'Supervisor verifies before action.',
          ],
          [
              'UC-03 FOD candidate',
              (
                  'Unrecognized object remains in a critical area beyond the'
                  ' configured duration.'
              ),
              'Creates a candidate alert and requests human verification.',
              'Supervisor validates clearance need.',
          ],
          [
              'UC-04 Alert review',
              'New alert enters the queue.',
              'Prioritizes, explains, and records acknowledgment.',
              'Faster, traceable decision.',
          ],
          [
              'UC-05 False-positive feedback',
              'User dismisses an incorrect alert.',
              'Captures reason and links to model/config version.',
              'Data supports later evaluation and tuning.',
          ],
      ],
  )

  add_header('8. AI, Data and Solution Requirements', level=2)
  add_header('8.1 Proposed prototype architecture', level=3)
  add_table(
      ['Layer', 'Prototype component', 'Responsibility'],
      [
          [
              'Experience',
              'Web supervisor console on Firebase Hosting or Cloud Run',
              'Alert queue, evidence view, status, feedback, dashboard.',
          ],
          [
              'Ingestion',
              'Cloud Run API / event endpoint',
              'Accept demo video events and metadata.',
          ],
          [
              'Vision',
              'Video/image analysis component',
              'Detect entities, zones, and candidate conditions.',
          ],
          [
              'Context',
              'Turnaround-state service',
              'Provide phase, gate configuration, and scenario rules.',
          ],
          [
              'Gen AI',
              'Gemini via an approved Google AI service',
              (
                  'Interpret multimodal evidence and produce'
                  ' schema-constrained explanation.'
              ),
          ],
          [
              'Orchestration',
              'Risk and alert service on Cloud Run',
              (
                  'Merge observations, context, deterministic guardrails, and'
                  ' priority.'
              ),
          ],
          [
              'Data',
              'Managed operational store and object storage',
              (
                  'Persist events, evidence references, configuration, and'
                  ' dispositions.'
              ),
          ],
          [
              'Analytics',
              'BigQuery or prototype reporting store',
              'Support evaluation metrics and trend summaries.',
          ],
          [
              'Security',
              'Identity, roles, logging, secrets',
              'Enforce access and provide auditability.',
          ],
      ],
  )

  add_header('8.2 Structured AI output contract', level=3)
  doc.add_paragraph(
      'The generative model response should be validated against a fixed schema'
      ' before it can be used by the alert service. Recommended fields:'
  )
  doc.add_paragraph('• scenario_type')
  doc.add_paragraph('• observation_summary')
  doc.add_paragraph('• evidence_references')
  doc.add_paragraph('• turnaround_context')
  doc.add_paragraph('• risk_factors')
  doc.add_paragraph('• uncertainty_or_missing_information')
  doc.add_paragraph('• recommended_human_check')
  doc.add_paragraph('• prohibited_action_flag')
  doc.add_paragraph('• model_version')

  add_quote(
      'Guardrail: The Gen AI component may explain and recommend a human check.'
      ' It must not issue binding operational clearance, identify a person,'
      ' infer intent, or claim regulatory compliance.'
  )

  add_header('8.3 Data entities', level=3)
  add_table(
      ['Entity', 'Minimum fields', 'Retention principle'],
      [
          [
              'Camera Feed',
              'camera_id, gate, viewpoint, authorization status',
              'Configuration record only.',
          ],
          [
              'Zone',
              'zone_id, geometry, scenario, effective version',
              'Versioned while referenced by events.',
          ],
          [
              'Turnaround context',
              'flight surrogate, gate, phase, effective time',
              'Minimize and retain only as needed.',
          ],
          [
              'Detection',
              'class, bounding area, confidence, timestamp',
              'Short-lived unless linked to event.',
          ],
          [
              'Safety event',
              'event_id, scenario, severity, context, reasoning, status',
              'Retained per approved policy.',
          ],
          [
              'Evidence',
              'frame/clip reference, capture time, integrity metadata',
              'Restricted and time-bounded.',
          ],
          [
              'Disposition',
              'user role, action, note, timestamp',
              'Audit-controlled.',
          ],
          [
              'Model/config version',
              'model, prompt, rule, threshold versions',
              'Retain for reproducibility.',
          ],
      ],
  )

  add_header('8.4 Training and evaluation data', level=3)
  doc.add_paragraph(
      '• Use synthetic, staged, public, or explicitly authorized footage for'
      ' the hackathon.'
  )
  doc.add_paragraph(
      '• Separate demonstration clips from evaluation clips to avoid'
      ' presenting memorized behavior as generalization.'
  )
  doc.add_paragraph(
      '• Represent relevant lighting, camera angle, occlusion, object scale,'
      ' and negative cases within the limited MVP scope.'
  )
  doc.add_paragraph(
      '• Document source, permission, labels, transformations, and known gaps'
      ' for each dataset.'
  )

  add_header('9. Non-Functional Requirements', level=2)
  add_table(
      ['ID', 'Category', 'Requirement / prototype target'],
      [
          [
              'NFR-01',
              'Performance',
              (
                  'For the curated demo, present an alert within 5 seconds of'
                  ' the candidate event entering the orchestration layer. Target'
                  ' requires validation.'
              ),
          ],
          [
              'NFR-02',
              'Availability',
              (
                  'Prototype shall recover gracefully from a failed model call'
                  ' and show a degraded-state message rather than fabricate an'
                  ' interpretation.'
              ),
          ],
          [
              'NFR-03',
              'Security',
              (
                  'Require authenticated access, role-based authorization,'
                  ' encrypted transport, managed secrets, and audit logging.'
              ),
          ],
          [
              'NFR-04',
              'Privacy',
              (
                  'Avoid facial recognition and identity inference; minimize'
                  ' person-level data; support configurable evidence retention.'
              ),
          ],
          [
              'NFR-05',
              'Accessibility',
              (
                  'Core console shall support keyboard navigation, visible'
                  ' focus, readable contrast, non-color severity labels, and'
                  ' concise language.'
              ),
          ],
          [
              'NFR-06',
              'Explainability',
              (
                  'Every alert shall show supporting evidence, context,'
                  ' confidence/uncertainty, and the rule/model version used.'
              ),
          ],
          [
              'NFR-07',
              'Scalability',
              (
                  'Services shall be modular so additional feeds and event'
                  ' types can be added without redesigning the user workflow.'
              ),
          ],
          [
              'NFR-08',
              'Reliability',
              (
                  'Duplicate events shall be correlated or suppressed based'
                  ' on configurable logic.'
              ),
          ],
          [
              'NFR-09',
              'Observability',
              (
                  'Capture service health, latency, model errors, alert counts,'
                  ' and user dispositions.'
              ),
          ],
          [
              'NFR-10',
              'Portability',
              (
                  'Deployment and configuration shall be reproducible from'
                  ' documented source and environment settings.'
              ),
          ],
          [
              'NFR-11',
              'Usability',
              (
                  'A supervisor shall be able to understand and disposition a'
                  ' demo alert in three interactions or fewer, excluding'
                  ' authentication.'
              ),
          ],
          [
              'NFR-12',
              'Language',
              'Submission materials and prototype UI shall be available in English.',
          ],
      ],
  )

  add_header('10. Responsible AI, Safety and Governance', level=2)
  add_header('10.1 Control model', level=3)
  add_table(
      ['Risk', 'Required control'],
      [
          [
              'False positive',
              (
                  'Confidence/uncertainty display, duplicate suppression,'
                  ' human verification, feedback capture.'
              ),
          ],
          [
              'False negative',
              (
                  'Explicit coverage limits, periodic evaluation, detector'
                  ' health monitoring, no claim of complete surveillance.'
              ),
          ],
          [
              'Hallucinated explanation',
              (
                  'Schema validation, evidence references, constrained prompt,'
                  ' deterministic rule checks, degraded-state handling.'
              ),
          ],
          [
              'Unsafe recommendation',
              (
                  'Approved recommendation catalogue; model cannot create'
                  ' operational commands or clearance.'
              ),
          ],
          [
              'Bias / uneven performance',
              (
                  'Evaluate across relevant conditions and equipment/person'
                  ' presentation; document gaps and avoid identity inference.'
              ),
          ],
          [
              'Privacy intrusion',
              (
                  'Authorized sources only, minimum evidence, access controls,'
                  ' retention limits, no facial recognition.'
              ),
          ],
          [
              'Overreliance',
              (
                  'Persistent decision-support labeling, training, and'
                  ' mandatory human disposition.'
              ),
          ],
          [
              'Security misuse',
              (
                  'Authentication, least privilege, secret management, audit'
                  ' logs, restricted evidence access.'
              ),
          ],
          [
              'Model drift',
              (
                  'Version tracking, periodic evaluation, threshold review,'
                  ' rollback capability.'
              ),
          ],
          [
              'Performance surveillance',
              (
                  'Do not use outputs to rank individual workers or infer'
                  ' attitude, intent, or productivity.'
              ),
          ],
      ],
  )

  add_header('10.2 Required user-facing notice', level=3)
  add_quote(
      'Decision-support notice: RampGuard AI provides safety observations and'
      ' recommendations for verification. It does not replace local procedures,'
      ' trained judgment, or authorized operational control.'
  )

  add_header('10.3 Governance checkpoints', level=3)
  doc.add_paragraph('• Data permission and privacy review completed before ingestion.')
  doc.add_paragraph(
      '• Scenario and recommended-response catalogue approved by a qualified'
      ' safety representative.'
  )
  doc.add_paragraph(
      '• Evaluation results reviewed before each model/configuration version'
      ' is promoted.'
  )
  doc.add_paragraph(
      '• Deployment has named owner, rollback path, and incident response'
      ' contact.'
  )
  doc.add_paragraph(
      '• Any move beyond prototype requires security, legal, privacy, safety,'
      ' operational, and regulatory review.'
  )

  add_header('11. Success Measures', level=2)
  doc.add_paragraph(
      'The following are proposed prototype targets, not current performance'
      ' claims. Final thresholds must be validated against the selected'
      ' scenario, dataset, and qualified stakeholder expectations.'
  )
  add_table(
      ['Metric', 'Definition', 'Proposed MVP target'],
      [
          [
              'Scenario recall',
              'Share of labeled in-scope events detected.',
              '≥ 85% on held-out curated evaluation clips.',
          ],
          [
              'Alert precision',
              'Share of presented alerts confirmed as relevant.',
              '≥ 80% on held-out curated evaluation clips.',
          ],
          [
              'Critical miss count',
              'Labeled critical demo events not presented.',
              '0 in the final scripted demo set.',
          ],
          [
              'Alert orchestration latency',
              'Time from candidate event to console alert.',
              '≤ 5 seconds at p95 in demo environment.',
          ],
          [
              'Explanation completeness',
              (
                  'Alerts containing evidence, context, uncertainty, and'
                  ' recommendation.'
              ),
              '100% of presented demo alerts.',
          ],
          [
              'Disposition traceability',
              'Closed events with complete decision history and versions.',
              '100%.',
          ],
          [
              'Usability',
              (
                  'Test users who can correctly explain the alert and next'
                  ' human check.'
              ),
              '≥ 80% in a small prototype usability exercise.',
          ],
          [
              'Demo reliability',
              'Successful end-to-end scripted demo runs.',
              '5 consecutive runs before submission.',
          ],
      ],
  )

  add_header('Impact measurement after prototype', level=3)
  doc.add_paragraph(
      '• Compare awareness time and information completeness against the'
      ' current or simulated baseline.'
  )
  doc.add_paragraph(
      '• Measure false-positive causes by scenario, camera condition, and'
      ' configuration version.'
  )
  doc.add_paragraph(
      '• Track whether alerts support learning and corrective action, without'
      ' using the system for individual performance ranking.'
  )

  add_header('12. MVP, Demo and Delivery Plan', level=2)
  add_header('12.1 MVP backlog by priority', level=3)
  add_table(
      ['Priority', 'Capability', 'Exit evidence'],
      [
          [
              'P0',
              (
                  'Reliable replay of curated gate video and candidate-event'
                  ' generation.'
              ),
              'Input produces timestamped event.',
          ],
          [
              'P0',
              'Zone/phase context and one critical scenario.',
              'Context changes event risk.',
          ],
          [
              'P0',
              'Gemini structured explanation with evidence and uncertainty.',
              'Valid schema rendered in UI.',
          ],
          [
              'P0',
              'Guardrailed severity and human disposition.',
              'Alert can be acknowledged, escalated, dismissed, resolved.',
          ],
          [
              'P0',
              'Cloud deployment.',
              'Public or evaluator-accessible deployed prototype as permitted.',
          ],
          [
              'P1',
              'Two additional scenarios and duplicate suppression.',
              'Scenario tests and cleaner alert queue.',
          ],
          [
              'P1',
              'Basic metrics dashboard.',
              'Evaluation and demo metrics visible.',
          ],
          [
              'P2',
              'Config admin and export.',
              'Scenario and event records manageable.',
          ],
      ],
  )

  add_header('12.2 Three-minute demonstration storyline', level=3)
  add_table(
      ['Time', 'Demo moment', 'Judge takeaway'],
      [
          [
              '0:00 - 0:20',
              (
                  'Show aircraft turnaround context and the specific safety'
                  ' problem.'
              ),
              'Clear problem alignment and user.',
          ],
          [
              '0:20 - 0:45',
              (
                  'Start a curated video and show configured safety zone /'
                  ' turnaround phase.'
              ),
              'Functional, understandable setup.',
          ],
          [
              '0:45 - 1:20',
              (
                  'Trigger the primary hazard; show detection, evidence,'
                  ' Gemini explanation, uncertainty, and risk guardrail.'
              ),
              'Meaningful Gen AI and sound implementation.',
          ],
          [
              '1:20 - 1:50',
              'Supervisor reviews and acknowledges or escalates the alert.',
              'Cohesive UX and human authority.',
          ],
          [
              '1:50 - 2:15',
              (
                  'Replay a visually similar condition in a safe phase and show'
                  ' suppression or lower severity.'
              ),
              'Contextual differentiation and reduced noise.',
          ],
          [
              '2:15 - 2:40',
              'Show audit timeline and evaluation dashboard.',
              'Traceability and measurable quality.',
          ],
          [
              '2:40 - 3:00',
              'Show scalable architecture and close with safety impact.',
              'Feasibility, scalability, and vision.',
          ],
      ],
  )

  add_header('12.3 Delivery workstreams', level=3)
  add_table(
      ['Workstream', 'Key outputs'],
      [
          [
              'Product & safety',
              (
                  'Scenario definitions, response catalogue, acceptance'
                  ' criteria, demo narrative.'
              ),
          ],
          [
              'Data & vision',
              (
                  'Approved clips, labels, zone detection, evaluation set,'
                  ' error analysis.'
              ),
          ],
          [
              'Gen AI & orchestration',
              (
                  'Prompt/schema, Gemini integration, context reasoning,'
                  ' guardrails, test harness.'
              ),
          ],
          [
              'Cloud & backend',
              (
                  'Cloud Run/Firebase deployment, event APIs, data store,'
                  ' logging, security.'
              ),
          ],
          [
              'UX',
              (
                  'Supervisor console, accessibility, states, acknowledgment'
                  ' and evidence views.'
              ),
          ],
          [
              'QA & submission',
              (
                  'Traceability, end-to-end tests, reliability runs, deck/PDF,'
                  ' three-minute video, documentation.'
              ),
          ],
      ],
  )

  add_header('13. Risks, Assumptions and Dependencies', level=2)
  add_header('13.1 Assumptions', level=3)
  doc.add_paragraph(
      '• The prototype will use recorded, simulated, synthetic, public, or'
      ' specifically approved video.'
  )
  doc.add_paragraph(
      '• A small set of observable scenarios can be defined with unambiguous'
      ' labels for demonstration.'
  )
  doc.add_paragraph(
      '• A qualified safety stakeholder will validate terminology and'
      ' recommended human checks before external presentation.'
  )
  doc.add_paragraph(
      '• The team can use an eligible Google AI model and deploy on an accepted'
      ' Google Cloud service.'
  )
  doc.add_paragraph('• Prototype targets are adjustable after baseline evaluation.')

  add_header('13.2 Dependencies', level=3)
  doc.add_paragraph(
      '• Access to Google Cloud project, billing/quota, identity, and approved'
      ' model endpoints.'
  )
  doc.add_paragraph('• Curated footage and labels with documented rights.')
  doc.add_paragraph(
      '• Scenario definitions, gate-zone mapping, and turnaround-state'
      ' simulation.'
  )
  doc.add_paragraph(
      '• Front-end, backend, AI, and testing capacity within the team.'
  )
  doc.add_paragraph(
      '• Published submission portal requirements and deadlines at time of'
      ' submission.'
  )

  add_header('13.3 Risk register', level=3)
  add_table(
      ['Risk', 'Likelihood', 'Impact', 'Mitigation'],
      [
          [
              'Insufficient or unlicensed aviation footage',
              'Medium',
              'High',
              'Use staged or synthetic clips and maintain provenance.',
          ],
          [
              'Computer vision unreliable under occlusion or low light',
              'High',
              'High',
              (
                  'Narrow scenarios, evaluate negative cases, display'
                  ' uncertainty, require human review.'
              ),
          ],
          [
              'Gen AI explanation not grounded in evidence',
              'Medium',
              'High',
              (
                  'Use structured evidence references, schema validation,'
                  ' deterministic checks, and fallback messages.'
              ),
          ],
          [
              'Too many alerts',
              'Medium',
              'High',
              (
                  'Context gating, temporal persistence, duplicate'
                  ' suppression, and threshold tuning.'
              ),
          ],
          [
              'Scope exceeds hackathon capacity',
              'High',
              'Medium',
              'Freeze P0 scope; prioritize one excellent scenario before adding more.',
          ],
          [
              'Cloud integration or quota issue',
              'Medium',
              'Medium',
              (
                  'Use modular services, local test doubles, monitoring, and'
                  ' a deterministic recorded demo path.'
              ),
          ],
          [
              'Safety claim overstated',
              'Medium',
              'High',
              (
                  'Label prototype and decision-support limits; report'
                  ' measured results only.'
              ),
          ],
          [
              'Sensitive data exposure',
              'Low/Medium',
              'High',
              (
                  'Approved data only, minimize retention, restrict access,'
                  ' remove identity features.'
              ),
          ],
      ],
  )

  add_header('14. Acceptance and Traceability', level=2)
  add_header('14.1 MVP acceptance criteria', level=3)
  add_table(
      ['ID', 'Acceptance criterion'],
      [
          [
              'AC-01',
              (
                  'Given an approved demo clip with an in-scope event, when the'
                  ' event occurs, then the system presents an alert with'
                  ' evidence, location/time, scenario, confidence/uncertainty,'
                  ' and explanation.'
              ),
          ],
          [
              'AC-02',
              (
                  'Given the same object pattern in a different configured'
                  ' turnaround phase, when context indicates lower risk, then'
                  ' the system suppresses or reduces priority according to a'
                  ' documented rule.'
              ),
          ],
          [
              'AC-03',
              (
                  'Given a model timeout or invalid structured response, when'
                  ' orchestration cannot validate the result, then the system'
                  ' displays a degraded-state message and does not fabricate'
                  ' an alert explanation.'
              ),
          ],
          [
              'AC-04',
              (
                  'Given an alert, when an authorized supervisor chooses a'
                  ' disposition, then the action, note, timestamp, and version'
                  ' metadata are stored and displayed in the event timeline.'
              ),
          ],
          [
              'AC-05',
              (
                  'Given an unauthorized user, when protected evidence or'
                  ' administrative functions are requested, then access is'
                  ' denied and logged.'
              ),
          ],
          [
              'AC-06',
              (
                  'Given keyboard-only navigation, when a user reviews the'
                  ' alert queue and dialog, then focus order and actionable'
                  ' controls are usable without a pointer.'
              ),
          ],
          [
              'AC-07',
              (
                  'Given the final build, when the scripted demonstration is'
                  ' run five consecutive times, then the P0 workflow completes'
                  ' successfully in every run.'
              ),
          ],
      ],
  )

  add_header('14.2 Requirement traceability matrix', level=3)
  add_table(
      ['Business need', 'Functional requirements', 'Acceptance evidence'],
      [
          [
              'Situational awareness',
              'FR-01 to FR-08',
              'AC-01, AC-02, scenario evaluation results.',
          ],
          [
              'Human-controlled response',
              'FR-08 to FR-10',
              'AC-04 and decision-support notice.',
          ],
          ['Safe failure', 'FR-06, FR-07', 'AC-03 and model failure test.'],
          [
              'Security and privacy',
              'FR-10, FR-14; NFR-03, NFR-04',
              'AC-05, access logs, retention configuration.',
          ],
          [
              'Usable experience',
              'FR-08, FR-09; NFR-05, NFR-11',
              'AC-06 and usability exercise.',
          ],
          [
              'Competition-ready prototype',
              'BR-02, FR-15',
              'AC-07, cloud deployment, demo video.',
          ],
      ],
  )

  add_header('Appendix A: Submission Checklist', level=2)
  doc.add_paragraph(
      '• Record a public three-minute demo using YouTube, Vimeo, or a public'
      ' Google Drive link, according to published instructions.'
  )
  doc.add_paragraph(
      '• Include clear architecture, data provenance, responsible-AI controls,'
      ' evaluation results, and limitations.'
  )
  doc.add_paragraph(
      '• Validate current rules, dates, and portal fields on the official site'
      ' before submission.'
  )

  add_header('Appendix B: Glossary', level=2)
  add_table(
      ['Term', 'Meaning'],
      [
          [
              'Apron / ramp',
              (
                  'Airport area used for aircraft parking, loading, unloading,'
                  ' servicing, and related ground activity.'
              ),
          ],
          [
              'FOD',
              (
                  'Foreign Object Debris, or an object that may create a'
                  ' hazard in the operating area.'
              ),
          ],
          [
              'Ground support equipment',
              'Vehicles and equipment used to service an aircraft on the ground.',
          ],
          [
              'Turnaround phase',
              (
                  'A defined stage in the sequence between aircraft arrival'
                  ' and departure.'
              ),
          ],
          [
              'Candidate event',
              (
                  'A machine-detected condition requiring contextual validation'
                  ' before presentation.'
              ),
          ],
          ['Disposition', 'The authorized user action taken on an alert.'],
          [
              'Gen AI',
              (
                  'Generative artificial intelligence used here for grounded'
                  ' interpretation and explanation.'
              ),
          ],
          [
              'Human in the loop',
              (
                  'A design in which a person retains review and decision'
                  ' authority.'
              ),
          ],
      ],
  )

  add_header('Appendix C: Sources and Basis', level=2)
  doc.add_paragraph(
      'Competition facts and mandatory requirements in this BRD are based on'
      ' the official AI Builder Cup 2026 Themes and FAQ pages accessed on 3'
      ' October 2026:'
  )
  doc.add_paragraph(
      '• AI Builder Cup 2026 - Themes, submission requirements, and judging'
      ' criteria: Themes | AI Builder Cup 2026'
  )
  doc.add_paragraph('• AI Builder Cup 2026 - FAQs: https://aibuildercup.com/Faqs.html')
  doc.add_paragraph()
  add_quote(
      'All RampGuard scope, architecture, requirements, targets, controls, and'
      ' demo recommendations are project proposals for hackathon planning. They'
      ' are not statements of an existing deployed capability, validated'
      ' airport procedure, or regulatory approval.'
  )

  doc.save('RampGuard_BRD.docx')


if __name__ == '__main__':
  create_brd_document()